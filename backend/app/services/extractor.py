"""Extract plain text from uploaded resume documents (PDF / DOCX / TXT)."""
from __future__ import annotations

import io

from pypdf import PdfReader

from app.core.config import settings
from app.core.logging import get_logger

logger = get_logger("extractor")

MAX_TEXT_CHARS = 60000


class UnsupportedFileError(Exception):
    pass


class EmptyResumeError(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        text = _extract_pdf(data)
    elif name.endswith(".docx"):
        text = _extract_docx(data)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedFileError(
            f"Unsupported file type. Allowed: {', '.join(settings.ALLOWED_EXTENSIONS)}"
        )

    text = clean_text(text)
    if not text:
        raise EmptyResumeError("No readable text was found in this resume. Try a text-based PDF.")
    return text[:MAX_TEXT_CHARS]


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise UnsupportedFileError("The PDF could not be read. It may be corrupted or password protected.")
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            continue
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError:  # pragma: no cover
        raise UnsupportedFileError("DOCX support is unavailable on the server.")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise UnsupportedFileError("The DOCX file could not be read.")
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def clean_text(text: str) -> str:
    import re

    text = text.replace("\x00", "")
    text = re.sub(r"[\r\n\t]+", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()